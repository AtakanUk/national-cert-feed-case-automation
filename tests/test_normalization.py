import io
from datetime import UTC, datetime
from unittest.mock import MagicMock

from openpyxl import load_workbook

from national_cert_feed_case_automation.excel import address_workbook
from national_cert_feed_case_automation.normalization import (
    build_incident,
    ip_artifacts,
    parse_notification,
    strip_html,
)
from national_cert_feed_case_automation.pipeline import Options, collect, repair_attachments

NOW = datetime(2026, 3, 1, 12, 0, tzinfo=UTC)

STRUCTURED = {
    "id": 501,
    "code": "VN-2026-501",
    "description": "<p>Exposed <b>RDP</b> service</p>",
    "explanation": "Remote desktop is reachable from the internet.",
    "action": "Restrict access.",
    "vulnerabilityAddresses": [
        {
            "id": 1,
            "address": "198.51.100.10",
            "port": 3389,
            "protocol_id": 1,
            "time_stamp": 1772366400000,
        },
        {
            "id": 2,
            "address": "198.51.100.10",
            "port": 3390,
            "protocol_id": 1,
            "time_stamp": 1772366400000,
        },
        {
            "id": 3,
            "address": "portal.example.com",
            "port": 443,
            "protocol_id": 3,
            "time_stamp": 1772366400000,
        },
    ],
    "files": [],
}
FILE_ONLY = {
    "id": 502,
    "code": "VN-2026-502",
    "vulnerabilityAddresses": [],
    "files": [{"id": 77, "name": "report.pdf"}],
}
TEXT_ONLY = {
    "id": 503,
    "code": "VN-2026-503",
    "vulnerabilityAddresses": [],
    "files": [],
    "address": "Hosts 203.0.113.5, 203.0.113.5 and 999.1.1.1",
}


def test_strip_html():
    assert strip_html("<b>Hello</b><br>World &amp; more") == "Hello World & more"


def test_parse_shapes():
    structured = parse_notification(STRUCTURED)
    assert structured.shape == "addresses" and structured.ips == ["198.51.100.10"]
    assert structured.addresses[2].protocol == "HTTP"
    assert parse_notification(FILE_ONLY).shape == "file"
    text = parse_notification(TEXT_ONLY)
    assert text.shape == "addresses" and text.ips == ["203.0.113.5"]


def test_build_incident_and_artifacts():
    notification = parse_notification(STRUCTURED)
    incident = build_incident(notification, "https://portal.example.com/#/case/{id}", [1019], NOW)
    assert incident["name"].endswith("VN-2026-501")
    assert "https://portal.example.com/#/case/501" in incident["description"]["content"]
    assert incident["properties"]["cert_ip_list"] == "198.51.100.10"
    assert "<td>3389</td>" in incident["properties"]["cert_address_table"]
    assert [a["value"] for a in ip_artifacts(notification)] == ["198.51.100.10"]


def test_address_workbook():
    rows = list(
        load_workbook(io.BytesIO(address_workbook(parse_notification(STRUCTURED)))).active.values
    )
    assert rows[0] == ("Record", "Seen (UTC)", "Address", "Port", "Protocol")
    assert rows[1][2:] == ("198.51.100.10", 3389, "TCP")
    assert list(
        load_workbook(io.BytesIO(address_workbook(parse_notification(TEXT_ONLY)))).active.values
    ) == [("Address",), ("203.0.113.5",)]


class FakeCert:
    base_url = "https://cert-api.example.com/v1"

    def __init__(self, notifications):
        self.notifications = {n["id"]: n for n in notifications}

    def notifications_since(self, since):
        return [{"id": nid} for nid in self.notifications]

    def notification(self, nid):
        if nid == 999:
            raise RuntimeError("HTTP 500")
        return self.notifications[nid]

    def file(self, file_id):
        return b"%PDF-1.4"


def test_collect_creates_incidents_with_evidence():
    soar = MagicMock()
    soar.cases_since.return_value = {500: {"id": 1}}
    soar.create_incident.side_effect = [{"id": 11}, {"id": 12}, {"id": 13}]
    cert = FakeCert([{"id": 500}, STRUCTURED, FILE_ONLY, TEXT_ONLY, {"id": 999}])

    result = collect(cert, soar, Options(), now=NOW)

    assert (result.seen, result.existing, result.created, result.failed) == (
        5,
        1,
        [11, 12, 13],
        [999],
    )
    attachments = [(call.args[0], call.args[1]) for call in soar.attach.call_args_list]
    assert attachments == [
        (11, "cert_VN-2026-501_addresses.xlsx"),
        (12, "report.pdf"),
        (13, "cert_VN-2026-503_addresses.xlsx"),
    ]
    assert soar.add_artifact.call_count == 2


def test_repair_only_touches_incidents_without_attachments():
    soar = MagicMock()
    soar.cases_since.return_value = {
        501: {"id": 21, "num_attachments": 0},
        502: {"id": 22, "num_attachments": 1},
    }
    assert repair_attachments(FakeCert([STRUCTURED, FILE_ONLY]), soar, Options()) == [21]
